"""Office document extraction: .docx, .xlsx, .pptx."""

from __future__ import annotations

from pathlib import Path

VERSION = "office-1.0"


def extract(path: Path, size_cap: int = 1_000_000) -> str:
    ext = path.suffix.lower()
    if ext == ".docx":
        text = _docx(path)
    elif ext == ".xlsx":
        text = _xlsx(path)
    elif ext == ".pptx":
        text = _pptx(path)
    else:
        raise ValueError(f"unsupported office format: {ext}")
    return text[:size_cap]


def _docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def _xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            parts.append(f"# Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append("\t".join(cells))
    finally:
        wb.close()
    return "\n".join(parts)


def _pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"# Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs)
                    if t.strip():
                        parts.append(t)
    return "\n".join(parts)
